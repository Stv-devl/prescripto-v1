"""Text extraction from PDF, DOCX, and XLSX files."""

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

_PAGE_NUMBER_RE = re.compile(
    r"^\s*(?:\d{1,3}\s*/\s*\d{0,3}|p(?:age)?\.?\s*\d{1,3})\s*$",
    re.IGNORECASE,
)


@dataclass
class ExtractedPage:
    """A page or section of extracted text."""

    text: str
    page: int
    source: str
    has_tables: bool = False


@dataclass
class ExtractionResult:
    """Full extraction result from a document."""

    pages: list[ExtractedPage] = field(default_factory=list)
    full_text: str = ""


def _table_to_markdown(table: fitz.table.Table) -> str:
    """Convert a PyMuPDF Table object to a markdown string with sentinel markers.

    Returns empty string for tables where < 30% of cells contain text,
    since these are likely decorative borders or layout grids with no real data.
    """
    rows = table.extract()
    if not rows:
        return ""

    cleaned: list[list[str]] = []
    for row in rows:
        cleaned.append([(cell or "").strip().replace("\n", " ") for cell in row])

    total_cells = sum(len(row) for row in cleaned)
    filled_cells = sum(1 for row in cleaned for cell in row if cell)
    if total_cells > 0 and filled_cells / total_cells < 0.3:
        logger.debug(
            "Skipping sparse table: %d/%d cells filled (%.0f%%)",
            filled_cells,
            total_cells,
            filled_cells / total_cells * 100,
        )
        return ""

    header = cleaned[0]
    if all(not h for h in header):
        header = [f"Col {j + 1}" for j in range(len(header))]
        data_rows = cleaned
    else:
        data_rows = cleaned[1:]

    if not data_rows:
        return ""

    last_col0 = ""
    for row in data_rows:
        if row and row[0]:
            last_col0 = row[0]
        elif row and not row[0] and last_col0:
            row[0] = last_col0

    lines: list[str] = []
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in data_rows:
        padded = row + [""] * (len(header) - len(row))
        lines.append("| " + " | ".join(padded[: len(header)]) + " |")

    md = "\n".join(lines)
    return f"\n\n<!-- TABLE -->\n{md}\n<!-- /TABLE -->\n\n"


def _blocks_overlap_table(
    block_bbox: tuple[float, float, float, float],
    table_bbox: tuple[float, float, float, float],
) -> bool:
    """Check if a text block significantly overlaps with a table bounding box.

    Returns True if > 50% of the block area is inside the table area.
    """
    x0 = max(block_bbox[0], table_bbox[0])
    y0 = max(block_bbox[1], table_bbox[1])
    x1 = min(block_bbox[2], table_bbox[2])
    y1 = min(block_bbox[3], table_bbox[3])

    if x1 <= x0 or y1 <= y0:
        return False

    intersection = (x1 - x0) * (y1 - y0)
    block_area = (block_bbox[2] - block_bbox[0]) * (block_bbox[3] - block_bbox[1])

    if block_area <= 0:
        return False

    return intersection / block_area > 0.5


def extract_pdf(file_path: Path) -> ExtractionResult:
    """Extract text from a PDF using PyMuPDF with structured table extraction.

    Tables are detected via find_tables() and converted to markdown with
    sentinel markers. Non-table text blocks are preserved in reading order.
    """
    doc = fitz.open(str(file_path))
    pages: list[ExtractedPage] = []
    all_text: list[str] = []

    for i, page in enumerate(doc):
        tables_result = page.find_tables()
        detected_tables = tables_result.tables if tables_result else []

        if not detected_tables:
            text = page.get_text("text").strip()
            if text:
                pages.append(ExtractedPage(text=text, page=i + 1, source=f"page {i + 1}"))
                all_text.append(text)
            continue

        table_markdowns: list[tuple[tuple[float, float, float, float], str]] = []
        valid_table_bboxes: list[tuple[float, float, float, float]] = []

        for table in detected_tables:
            md = _table_to_markdown(table)
            bbox = table.bbox
            table_markdowns.append((bbox, md))
            if md:
                valid_table_bboxes.append(bbox)

        page_dict = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
        blocks = page_dict.get("blocks", [])

        elements: list[tuple[float, str, str]] = []

        for block in blocks:
            if block.get("type") != 0:
                continue
            block_bbox = (block["bbox"][0], block["bbox"][1], block["bbox"][2], block["bbox"][3])

            overlaps_valid = any(_blocks_overlap_table(block_bbox, tb) for tb in valid_table_bboxes)
            if overlaps_valid:
                continue

            block_text_parts: list[str] = []
            for line in block.get("lines", []):
                line_text = "".join(span.get("text", "") for span in line.get("spans", []))
                if line_text.strip():
                    block_text_parts.append(line_text.strip())

            if block_text_parts:
                block_text = "\n".join(block_text_parts)
                if _PAGE_NUMBER_RE.match(block_text):
                    continue
                elements.append((block_bbox[1], "text", block_text))

        for bbox, md in table_markdowns:
            if md:
                elements.append((bbox[1], "table", md))

        elements.sort(key=lambda e: e[0])

        page_parts = [content for _, _, content in elements]
        text = "\n\n".join(part.strip() for part in page_parts if part.strip())

        if text:
            has_tables = any(etype == "table" for _, etype, _ in elements)
            pages.append(
                ExtractedPage(text=text, page=i + 1, source=f"page {i + 1}", has_tables=has_tables)
            )
            all_text.append(text)
            logger.debug(
                "Page %d: %d tables detected, %d text blocks kept",
                i + 1,
                len(detected_tables),
                sum(1 for _, t, _ in elements if t == "text"),
            )

    doc.close()
    return ExtractionResult(pages=pages, full_text="\n\n".join(all_text))


def extract_docx(file_path: Path) -> ExtractionResult:
    """Extract text from a DOCX preserving paragraph structure."""
    doc = DocxDocument(str(file_path))
    paragraphs: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    full_text = "\n\n".join(paragraphs)
    pages = [ExtractedPage(text=full_text, page=1, source="document")] if full_text else []
    return ExtractionResult(pages=pages, full_text=full_text)


def extract_xlsx(file_path: Path) -> ExtractionResult:
    """Extract text from an XLSX, one page per sheet."""
    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    pages: list[ExtractedPage] = []
    all_text: list[str] = []

    for idx, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))

        if rows:
            text = "\n".join(rows)
            pages.append(ExtractedPage(text=text, page=idx + 1, source=f"sheet '{sheet_name}'"))
            all_text.append(f"[{sheet_name}]\n{text}")

    wb.close()
    return ExtractionResult(pages=pages, full_text="\n\n".join(all_text))


def extract(file_path: Path) -> ExtractionResult:
    """Extract text from a file based on its extension."""
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(file_path)
    elif ext == ".docx":
        return extract_docx(file_path)
    elif ext == ".xlsx":
        return extract_xlsx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
